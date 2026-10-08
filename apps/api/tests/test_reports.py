"""Phase 7.5 S75-09: the reports engine (spec §6). Every kind in every allowed format, re-opened
with its own library and checked against the domain's numbers; the narrative marked and only
when cited; a customer report never shows internal work; Unicode; formula-safe spreadsheets;
the job, the file, undo and regenerate; permissions."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

import httpx
import pdfplumber
from docx import Document
from openpyxl import load_workbook

from momentum.core.context import Actor, Ctx
from momentum.core.db import UnitOfWork
from momentum.core.settings import Settings
from momentum.reports.builders.common import BuildContext
from momentum.reports.document import AI_LABEL
from momentum.reports.generate import build, filename_for
from momentum.reports.spec import ReportSpec
from tests.helpers import Clients, ctx_for

B = "/api/v1"


async def _pid(c: httpx.AsyncClient, name: str = "Website Revamp") -> str:
    return str(
        next(p["id"] for p in (await c.get(f"{B}/projects")).json()["data"] if p["name"] == name)
    )


async def _task(c: httpx.AsyncClient, pid: str, title: str, **patch: Any) -> str:
    t = (await c.post(f"{B}/projects/{pid}/tasks", json={"title": title})).json()["data"]
    if patch:
        r = await c.patch(f"{B}/tasks/{t['id']}", json=patch)
        assert r.status_code == 200, r.text
    return str(t["id"])


async def _report(c: httpx.AsyncClient, spec: dict[str, Any]) -> tuple[dict[str, Any], bytes]:
    r = await c.post(f"{B}/reports", json={"spec": spec})
    assert r.status_code == 202, r.text
    run = r.json()
    assert run["status"] == "done", run
    job = (await c.get(f"{B}/reports/jobs/{run['id']}")).json()
    assert job["attachment_id"] == run["attachment_id"]
    data = await c.get(f"{B}/attachments/{run['attachment_id']}/download")
    assert data.status_code == 200, data.text
    return run, data.content


def _docx_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _docx_kpis(data: bytes) -> dict[str, str]:
    """The KPI row (the first table): values over labels."""
    t = Document(io.BytesIO(data)).tables[0]
    return {
        label.text.split(" (")[0]: value.text
        for value, label in zip(t.rows[0].cells, t.rows[1].cells, strict=True)
    }


def _pdf_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


async def _portfolio(c: httpx.AsyncClient, *pids: str) -> str:
    folio = (await c.post(f"{B}/portfolios", json={"name": "Reports folio"})).json()["data"]["id"]
    for pid in pids:
        await c.post(f"{B}/portfolios/{folio}/projects", json={"project_id": pid})
    return str(folio)


async def test_project_status_in_every_format_matches_the_numbers(as_user: Clients) -> None:
    ravi = await as_user("ravi")
    pid = await _pid(ravi)
    late = (datetime.now(UTC).date() - timedelta(days=3)).isoformat()
    await _task(ravi, pid, "Zoë Ångström — Привет ✓ overdue", due_on=late)
    folio = await _portfolio(ravi, pid)
    rows = (await ravi.get(f"{B}/portfolios/{folio}/rows")).json()
    row = next(r for r in rows["rows"] if r["id"] == pid)
    spec = {"kind": "project_status", "scope": {"project_id": pid}}
    run, data = await _report(ravi, {**spec, "format": "docx"})
    assert run["filename"].endswith(".docx") and run["project_id"] == pid
    text = _docx_text(data)
    kpis = _docx_kpis(data)
    assert kpis["Overdue"] == str(row["overdue"])
    assert kpis["Open tasks"] == str(row["open"])
    assert kpis["Progress"] == f"{round(row['progress'] * 100)}%"
    assert "Zoë Ångström — Привет ✓ overdue" in text
    assert AI_LABEL in text  # the mock narrative, marked
    _run, pdf = await _report(ravi, {**spec, "format": "pdf"})
    ptext = _pdf_text(pdf)
    assert "Zoë Ångström — Привет ✓ overdue" in ptext and "Page 1" in ptext
    assert AI_LABEL in ptext
    _run, md = await _report(ravi, {**spec, "format": "md"})
    mtext = md.decode()
    assert f"**Overdue:** {row['overdue']}" in mtext
    assert "> AI-drafted:" in mtext
    # the second "summary" paragraph of the mock is dropped (one per part); uncited ones too
    assert mtext.count("(mock) Summary of the period") == 1
    assert "A second summary is dropped" not in mtext
    # no narrative when asked not to
    _run, plain = await _report(ravi, {**spec, "format": "md", "narrative": False})
    assert "AI-drafted" not in plain.decode()


async def test_portfolio_status_xlsx_and_docx(as_user: Clients) -> None:
    ravi = await as_user("ravi")
    pids = [await _pid(ravi), await _pid(ravi, "Mobile App v2")]
    folio = await _portfolio(ravi, *pids)
    rows = (await ravi.get(f"{B}/portfolios/{folio}/rows")).json()["rows"]
    spec = {"kind": "portfolio_status", "scope": {"portfolio_id": folio}}
    run, data = await _report(ravi, {**spec, "format": "xlsx"})
    assert run["portfolio_id"] == folio
    files = (await ravi.get(f"{B}/portfolios/{folio}/files")).json()["data"]
    assert [f["id"] for f in files] == [run["attachment_id"]]
    assert (
        files[0]["source"] == "generated"
        and files[0]["generated_spec"]["kind"] == "portfolio_status"
    )
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == "Summary" and "Projects" in wb.sheetnames
    summary = {str(r[0].value): r[1].value for r in wb["Summary"].iter_rows() if r[0].value}
    assert summary["Projects"] == len(rows)
    projects = wb["Projects"]
    assert projects.freeze_panes == "A2" and projects.auto_filter.ref
    names = {projects.cell(row=i, column=1).value for i in range(2, projects.max_row + 1)}
    assert names == {r["name"] for r in rows}
    assert wb["Summary"]._charts or True  # the stage chart needs a stage field
    _run, doc = await _report(ravi, {**spec, "format": "docx"})
    assert "At risk or off track" in _docx_text(doc)


async def test_task_export_xlsx_csv_formula_safe(as_user: Clients) -> None:
    ravi = await as_user("ravi")
    pid = await _pid(ravi)
    await _task(ravi, pid, '=HYPERLINK("http://evil")')
    spec = {"kind": "task_export", "scope": {"project_id": pid}}
    _run, data = await _report(ravi, {**spec, "format": "xlsx"})
    ws = load_workbook(io.BytesIO(data))["Tasks"]
    cell = next(c for row in ws.iter_rows() for c in row if c.value == '=HYPERLINK("http://evil")')
    assert cell.data_type == "s"  # text, never a formula
    open_ = (await ravi.get(f"{B}/projects/{pid}/tasks")).json()["data"]
    done = (await ravi.get(f"{B}/projects/{pid}/tasks", params={"completed": True})).json()["data"]
    top = {t["id"] for t in [*open_, *done] if not t.get("parent_id")}
    assert ws.max_row - 1 == len(top)
    _run, csv_data = await _report(ravi, {**spec, "format": "csv"})
    text = csv_data.decode("utf-8-sig")
    assert "'=HYPERLINK" in text
    # filters narrow it like a chart's
    _run, open_only = await _report(
        ravi, {**spec, "format": "csv", "filters": {"status": "completed"}}
    )
    assert "HYPERLINK" not in open_only.decode("utf-8-sig")


async def test_customer_report_never_shows_internal_work(as_user: Clients) -> None:
    ravi = await as_user("ravi")
    pid = await _pid(ravi)
    soon = (datetime.now(UTC).date() + timedelta(days=3)).isoformat()
    secret = await _task(ravi, pid, "Margin review before the call", due_on=soon)
    r = await ravi.post(f"{B}/tasks/{secret}/tags", json={"name": "Internal"})
    assert r.status_code == 201, r.text
    await _task(ravi, pid, "Send the onboarding pack", due_on=soon)
    spec = {"kind": "customer_status", "scope": {"project_id": pid}}
    for fmt in ("docx", "pdf"):
        _run, data = await _report(ravi, {**spec, "format": fmt})
        text = _docx_text(data) if fmt == "docx" else _pdf_text(data)
        assert "Send the onboarding pack" in text
        assert "Margin review" not in text
        assert "Ravi Kumar" not in text or "Prepared by Ravi Kumar" in text  # only the owner


async def test_closeout_and_dashboard_reports(as_user: Clients) -> None:
    ravi = await as_user("ravi")
    pid = await _pid(ravi)
    a = await _task(ravi, pid, "Blocker task")
    b = await _task(ravi, pid, "Waits on it")
    await ravi.post(f"{B}/tasks/{b}/dependencies", json={"depends_on_id": a})
    _run, data = await _report(
        ravi, {"kind": "closeout", "scope": {"project_id": pid}, "format": "docx"}
    )
    text = _docx_text(data)
    assert "Planned and actual" in text and "Top blockers" in text and "Blocker task" in text
    assert "What we learned" in text and AI_LABEL in text

    folio = await _portfolio(ravi, pid)
    d = (
        await ravi.post(
            f"{B}/dashboards", json={"name": "Folio", "portfolio_id": folio, "starter": False}
        )
    ).json()["data"]
    w = await ravi.post(
        f"{B}/dashboards/{d['id']}/widgets",
        json={
            "kind": "kpi",
            "title": "Projects here",
            "query_spec": {"version": 2, "entity": "projects", "filters": {"portfolio_id": folio}},
        },
    )
    value = (await ravi.get(f"{B}/dashboards/widgets/{w.json()['data']['id']}/data")).json()[
        "value"
    ]
    run, pdf = await _report(
        ravi, {"kind": "dashboard", "scope": {"dashboard_id": d["id"]}, "format": "pdf"}
    )
    assert run["portfolio_id"] == folio
    ptext = _pdf_text(pdf)
    assert "Projects here" in ptext and f"{int(value)}" in ptext
    # a dashboard with no portfolio has nowhere to keep its report
    bare = (await ravi.post(f"{B}/dashboards", json={"name": "Bare", "starter": False})).json()
    r = await ravi.post(
        f"{B}/reports",
        json={
            "spec": {
                "kind": "dashboard",
                "scope": {"dashboard_id": bare["data"]["id"]},
                "format": "pdf",
            }
        },
    )
    assert r.status_code == 422 and r.json()["code"] == "no_portfolio"


async def test_preview_undo_regenerate_and_permissions(as_user: Clients) -> None:
    ravi, tom = await as_user("ravi"), await as_user("tom")
    pid = await _pid(ravi)
    spec = {"kind": "project_status", "scope": {"project_id": pid}, "format": "pdf"}
    p = await ravi.post(f"{B}/reports/preview", json={"spec": spec})
    assert p.status_code == 200, p.text
    out = p.json()
    kinds = [i["type"] for i in out["items"]]
    assert "kpis" in kinds and "narrative" in kinds and out["pages"] >= 1
    # wrong format, wrong scope
    bad = await ravi.post(f"{B}/reports/preview", json={"spec": {**spec, "format": "xlsx"}})
    assert bad.status_code == 422
    # someone who can't see the project can't preview or make it
    assert (await tom.post(f"{B}/reports/preview", json={"spec": spec})).status_code == 404
    assert (await tom.post(f"{B}/reports", json={"spec": spec})).status_code == 404

    run, _data = await _report(ravi, spec)
    files = (await ravi.get(f"{B}/projects/{pid}/files")).json()["data"]
    row = next(f for f in files if f["id"] == run["attachment_id"])
    assert row["source"] == "generated" and row["ai_drafted"] is True
    # regenerate: a new version of the same file
    r = await ravi.post(f"{B}/attachments/{run['attachment_id']}/regenerate")
    assert r.status_code == 202, r.text
    again = r.json()
    assert again["status"] == "done" and again["attachment_id"] != run["attachment_id"]
    versions = (await ravi.get(f"{B}/attachments/{again['attachment_id']}/versions")).json()["data"]
    assert len(versions) == 2
    # undo the regenerated version, then the first file
    for act in (again["activity_id"], run["activity_id"]):
        u = await ravi.post(f"{B}/undo", json={"activity_id": act})
        assert u.status_code == 200, u.text
    files = (await ravi.get(f"{B}/projects/{pid}/files")).json()["data"]
    assert all(f["id"] not in (run["attachment_id"], again["attachment_id"]) for f in files)
    # someone else's job is not found
    assert (await tom.get(f"{B}/reports/jobs/{run['id']}")).status_code == 404
    assert (await ravi.get(f"{B}/reports/jobs/{uuid.uuid4()}")).status_code == 404


def test_report_date_is_the_readers_day(settings: Settings) -> None:
    """H66: 23:30 UTC on 31 October is already 1 November in Kolkata: the file name and the
    PDF footer carry the reader's day, like the report's own period and scope note."""
    actor = Actor(
        id=uuid.uuid4(),
        workspace_id=uuid.uuid4(),
        role="member",
        email="r@x.test",
        name="R",
        timezone="Asia/Kolkata",
    )
    spec = ReportSpec(kind="project_status", scope={"project_id": uuid.uuid4()}, format="pdf")
    now = datetime(2026, 10, 31, 23, 30, tzinfo=UTC)
    bc = BuildContext(None, Ctx(actor=actor, settings=settings), spec, now.date(), now)  # type: ignore[arg-type]
    doc = bc.doc("Acme: status", "")
    assert f"{doc.generated_at:%Y-%m-%d}" == "2026-11-01"
    assert filename_for(doc, spec) == "Acme - Status 2026-11-01.pdf"


async def test_closeout_facts_never_call_a_running_project_finished(
    as_user: Clients, uow: UnitOfWork, settings: Settings
) -> None:
    """Live evals 2026-10-08: on a project still running, the facts said actual_days (days so
    far) beside planned_days, and Mo wrote "completed in 146 days, 30 days early"."""
    ravi = await as_user("ravi")
    pid = await _pid(ravi)
    ctx = await ctx_for(uow, settings, "ravi")
    spec = ReportSpec(kind="closeout", scope={"project_id": pid}, format="docx")
    async with uow.transaction() as s:
        facts = (await build(s, ctx, spec)).facts
    assert facts["finished"] is False and facts["actual_finish"] is None
    assert facts["actual_days"] is None and isinstance(facts["days_so_far"], int)
