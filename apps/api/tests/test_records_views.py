"""Phase 7.6 S76-04 (spec §6.6-§6.7): records on dashboards (entities ``records`` and
``record_lines``, money per currency, empty states), the "Accounts payable" role template, the
``records_export`` report (xlsx and csv), Mo's records tools, and the guest rule by every route:
a guest never sees a financial record through the API, search, Mo, a dashboard or an export."""

from __future__ import annotations

import io
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import update
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from momentum.core.db import UnitOfWork
from momentum.domain.dashboards.query_v2 import drill_any, run_any
from momentum.domain.dashboards.role_templates import catalog, load
from momentum.domain.dashboards.schemas import DrillAnyIn
from momentum.domain.dashboards.schemas_v2 import RecordsSpec
from momentum.domain.records import service as records
from momentum.domain.users.models import User
from momentum.reports.builders.common import BuildContext
from momentum.reports.generate import build
from momentum.reports.render.text import render_csv
from momentum.reports.render.xlsx import render_xlsx
from momentum.reports.spec import ReportSpec
from tests.ai_fixtures import World, call, world
from tests.helpers import ctx_for
from tests.jobs_env import JobsEnv
from tests.test_records import _agent_ctx, bill

_ = world, BuildContext


@pytest.fixture
def make_env(
    uow: UnitOfWork,
    session_factory: async_sessionmaker[AsyncSession],
    tmp_path: Path,
    world: World,
) -> Callable[..., JobsEnv]:
    return lambda **kw: JobsEnv(uow, session_factory, tmp_path, world, **kw)


async def _seed(make_env: Callable[..., JobsEnv], world: World) -> JobsEnv:
    env = make_env()
    await env.install("echo")
    impl = env.packs.packs["echo"].record_type("echo_bill")
    agent = await _agent_ctx(env)
    project_id, task_id = world.project.id, world.copy.id
    for over in (
        {},
        {"number": "2", "total": "50.00", "lines": [], "dated": "2026-10-02"},
        {"number": "3", "total": "70.00", "currency": "EUR", "lines": []},
        {"vendor": {"name": "Globex"}, "number": "4", "total": "10.00", "lines": [], "status": "x"},
    ):
        over.pop("status", None)
        async with env.uow.transaction() as s:
            await records.create_record(
                s, agent, impl, project_id=project_id, task_id=task_id, data=bill(**over)
            )
    return env


def _spec(**over: Any) -> RecordsSpec:
    base: dict[str, Any] = {"version": 2, "entity": "records", "type": "echo_bill"}
    base.update(over)
    return RecordsSpec.model_validate(base)


async def _guest(env: JobsEnv, world: World) -> Any:
    async with env.uow.transaction() as s:
        await s.execute(update(User).where(User.id == world.lena.actor.id).values(role="guest"))
    return await ctx_for(env.uow, env.settings, "lena")


async def test_dashboard_widgets_over_records(
    make_env: Callable[..., JobsEnv], world: World
) -> None:
    env = await _seed(make_env, world)
    async with env.uow.transaction() as s:
        bar = await run_any(
            s,
            world.ravi,
            "bar",
            _spec(group_by="vendor.name", measure="sum", measure_path="total", window_days=730),
        )
        count = await run_any(s, world.ravi, "count", _spec(window_days=730))
        money_kpi = await run_any(
            s, world.ravi, "kpi", _spec(measure="sum", measure_path="total", window_days=730)
        )
        lines = await run_any(
            s,
            world.ravi,
            "donut",
            _spec(
                entity="record_lines",
                array="lines",
                group_by="lines[].description",
                measure="sum",
                measure_path="lines[].amount",
                window_days=730,
            ),
        )
        missing = await run_any(s, world.ravi, "count", _spec(type="receipt"))  # no pack has it
        drill = await drill_any(
            s,
            world.ravi,
            DrillAnyIn(
                query_spec=_spec(group_by="vendor.name", measure="sum", measure_path="total"),
                key="Acme Ltd · USD",
            ),
        )
    assert {g.label: g.value for g in bar.groups} == {
        "Acme Ltd · USD": 150.0,
        "Acme Ltd · EUR": 70.0,
        "Globex · USD": 10.0,
    }
    assert count.value == 4.0
    # a KPI over money in two currencies never adds them: one group per currency instead
    assert money_kpi.value is None
    assert {g.label: g.value for g in money_kpi.groups} == {"USD": 160.0, "EUR": 70.0}
    assert {g.label: g.value for g in lines.groups} == {
        "Widgets · USD": 60.0,
        "Gadgets · USD": 40.0,
    }
    assert missing.value is None and missing.notes == ["No receipt records yet"]
    assert drill.entity == "records" and drill.total == 2
    assert {r["title"] for r in drill.records} == {"Acme Ltd INV-0041", "Acme Ltd 2"}


def test_the_accounts_payable_template() -> None:
    names = {t["key"]: t for t in catalog()}
    ap = names["accounts_payable"]
    assert ap["name"] == "Accounts payable" and len(ap["widgets"]) == 6
    assert all(w["spec"]["type"] == "invoice" for w in load("accounts_payable")["widgets"])


async def test_accounts_payable_widgets_are_empty_states_before_any_invoice(
    make_env: Callable[..., JobsEnv], world: World
) -> None:
    from momentum.domain.dashboards.role_templates import SPEC

    env = make_env()
    async with env.uow.transaction() as s:
        for w in load("accounts_payable")["widgets"]:
            out = await run_any(s, world.ravi, w["kind"], SPEC.validate_python(w["spec"]))
            assert out.notes == ["No invoice records yet"], w["title"]


async def test_records_export(make_env: Callable[..., JobsEnv], world: World) -> None:
    env = await _seed(make_env, world)
    spec = ReportSpec(
        kind="records_export",
        scope={"project_id": world.project.id},  # type: ignore[arg-type]
        format="xlsx",
        record_type="echo_bill",
    )
    async with env.uow.transaction() as s:
        doc = await build(s, world.ravi, spec)
    wb = load_workbook(io.BytesIO(render_xlsx(doc)))
    assert {"Records", "Lines"} <= set(wb.sheetnames)
    sheet = wb["Records"]
    head = [c.value for c in sheet[1]]
    assert head[:2] == ["Title", "Status"] and head[-1] == "Currency"
    totals = {
        (row[0].value, row[-1].value): row[head.index("Total")].value
        for row in sheet.iter_rows(min_row=2)
    }
    assert totals[("Acme Ltd 3", "EUR")] == 70.0  # money as a number, its currency beside it
    lines = wb["Lines"]
    assert [c.value for c in lines[1]][:2] == ["Record", "#"]
    assert lines.max_row == 3  # two items of the one bill with lines
    csv_text = render_csv(doc).decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("Title,Status")
    with pytest.raises(ValueError, match="names its record type"):
        ReportSpec(kind="records_export", scope={"project_id": world.project.id}, format="csv")  # type: ignore[arg-type]


async def test_mo_reads_records_with_server_numbers(
    make_env: Callable[..., JobsEnv], world: World
) -> None:
    env = await _seed(make_env, world)
    found = await call(
        env.uow,
        world.ravi,
        "search_records",
        {"type": "echo_bill", "text": "Globex"},
        mode="execute",
    )
    assert found.ok and [r["title"] for r in found.result.data["records"]] == ["Globex 4"]
    rid = found.result.data["records"][0]["id"]
    one = await call(env.uow, world.ravi, "get_record", {"id": rid}, mode="execute")
    assert one.ok and one.result.data["data"]["number"] == "4"
    summed = await call(
        env.uow,
        world.ravi,
        "query_records",
        {
            "type": "echo_bill",
            "group_by": ["vendor.name"],
            "measures": [{"op": "sum", "path": "total"}],
        },
        mode="execute",
    )
    rows = {
        (r["group"]["vendor.name"], r["currency"]): r["values"]["sum(total)"]
        for r in summed.result.data["rows"]
    }
    assert rows == {("Acme Ltd", "USD"): 150.0, ("Acme Ltd", "EUR"): 70.0, ("Globex", "USD"): 10.0}


async def test_a_guest_never_sees_a_financial_record_by_any_route(
    make_env: Callable[..., JobsEnv], world: World
) -> None:
    env = await _seed(make_env, world)
    guest = await _guest(env, world)
    # Mo
    found = await call(env.uow, guest, "search_records", {"type": "echo_bill"}, mode="execute")
    assert found.result.data["records"] == []
    summed = await call(env.uow, guest, "query_records", {"type": "echo_bill"}, mode="execute")
    assert summed.result.data["rows"] == []
    # dashboards
    async with env.uow.transaction() as s:
        out = await run_any(s, guest, "count", _spec(window_days=730))
    assert out.value in (None, 0.0) and out.tasks_total == 0
    # exports: a guest can't add a report to the project, and the builder shows nothing anyway
    spec = ReportSpec(
        kind="records_export",
        scope={"project_id": world.project.id},  # type: ignore[arg-type]
        format="csv",
        record_type="echo_bill",
    )
    async with env.uow.transaction() as s:
        doc = await build(s, guest, spec)
    assert render_csv(doc).decode("utf-8-sig").strip().splitlines() == [
        "Title,Status,Vendor name,Number,Dated,Total,Currency"
    ]
    _ = datetime.now(UTC)
