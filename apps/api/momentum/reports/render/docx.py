"""Phase 7.5 (spec §6.1): the Word renderer (python-docx). Styled headings, tables whose header row
repeats on every page, charts as pictures, and every AI paragraph behind an amber left rule with
"AI-drafted, review before sending"."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from momentum.reports.document import (
    AI_LABEL,
    Callout,
    Chart,
    Heading,
    KPIRow,
    PageBreak,
    Paragraph,
    ReportDocument,
    Table,
    TaskList,
)
from momentum.reports.render.png import render_png
from momentum.reports.render.text import text_of

AMBER = "C98500"
TONES = {"info": "2A78D6", "warn": "C98500", "crit": "E34948", "ok": "1BAF7A"}
FONT = "DejaVu Sans"


def _left_rule(par: object, color: str) -> None:
    p_pr = par._p.get_or_add_pPr()  # type: ignore[attr-defined]
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    p_pr.append(borders)


def _repeat_header(row: object) -> None:
    tr_pr = row._tr.get_or_add_trPr()  # type: ignore[attr-defined]
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def render_docx(doc: ReportDocument) -> bytes:
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
    d.add_heading(doc.title, level=0)
    sub = d.add_paragraph(doc.subtitle)
    sub.runs[0].italic = True
    note = d.add_paragraph(doc.scope_note)
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = RGBColor(0x6E, 0x74, 0x7E)
    for b in doc.blocks:
        if isinstance(b, Heading):
            d.add_heading(b.text, level=min(3, b.level))
        elif isinstance(b, Paragraph):
            if b.ai:
                label = d.add_paragraph()
                run = label.add_run(AI_LABEL)
                run.bold = True
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string(AMBER)
                _left_rule(label, AMBER)
                par = d.add_paragraph(b.text)
                _left_rule(par, AMBER)
            else:
                d.add_paragraph(b.text)
        elif isinstance(b, KPIRow):
            t = d.add_table(rows=2, cols=len(b.items))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, k in enumerate(b.items):
                c = t.cell(0, i).paragraphs[0].add_run(k.value)
                c.bold = True
                c.font.size = Pt(16)
                lbl = (
                    t.cell(1, i)
                    .paragraphs[0]
                    .add_run(k.label + (f" ({k.delta})" if k.delta else ""))
                )
                lbl.font.size = Pt(8.5)
        elif isinstance(b, Table):
            d.add_paragraph().add_run(b.title).bold = True
            if not b.rows:
                d.add_paragraph("Nothing here.")
                continue
            t = d.add_table(rows=1, cols=len(b.columns))
            t.style = "Light Grid Accent 1"
            for i, c in enumerate(b.columns):
                t.rows[0].cells[i].text = c
            _repeat_header(t.rows[0])
            for r in b.rows:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = text_of(v)
        elif isinstance(b, TaskList):
            d.add_paragraph().add_run(b.title).bold = True
            if not b.items:
                d.add_paragraph(b.empty)
            for it in b.items:
                bits = [
                    x for x in (it.assignee, f"due {it.due_on:%d %b}" if it.due_on else None) if x
                ]
                d.add_paragraph(
                    f"{'✓ ' if it.done else ''}{it.key} {it.title}"
                    + (f" ({', '.join(bits)})" if bits else ""),
                    style="List Bullet",
                )
        elif isinstance(b, Callout):
            par = d.add_paragraph(b.text)
            _left_rule(par, TONES[b.tone])
        elif isinstance(b, Chart):
            if not b.values:
                continue
            d.add_paragraph().add_run(b.title).bold = True
            d.add_picture(io.BytesIO(render_png(b)), width=Cm(15))
        elif isinstance(b, PageBreak):
            d.add_page_break()  # type: ignore[no-untyped-call]
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
