"""Phase 7.5 (spec §4.6, §4.7): pixels for the vision model, only when a person asks Mo to look.

A PDF page is rendered with pypdfium2 at 150 DPI; an image, or a picture embedded in a document
or slide, is decoded with Pillow. Either way the result is re-encoded as a JPEG (quality 85) with
the long edge at most 1568 px, which drops EXIF and GPS data (nothing but pixels is written). The
40-megapixel limit is checked from the header before decoding.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from momentum.files.safety import MAX_IMAGE_PIXELS, UnsafeFile, check_image_header

DPI = 150
LONG_EDGE = 1568
QUALITY = 85


@dataclass(frozen=True)
class Rendered:
    jpeg: bytes
    width: int
    height: int
    locator: str


def _encode(im: object, locator: str) -> Rendered:
    from PIL import Image

    assert isinstance(im, Image.Image)
    if im.mode not in ("RGB", "L"):
        background = Image.new("RGB", im.size, (255, 255, 255))
        rgba = im.convert("RGBA")
        background.paste(rgba, mask=rgba.split()[-1])
        im = background
    w, h = im.size
    scale = min(1.0, LONG_EDGE / max(w, h))
    if scale < 1.0:
        im = im.resize(
            (max(1, round(w * scale)), max(1, round(h * scale))), Image.Resampling.LANCZOS
        )
    out = io.BytesIO()
    # a fresh JPEG with no exif/icc arguments: metadata from the source never survives
    im.save(out, format="JPEG", quality=QUALITY, optimize=True)
    return Rendered(out.getvalue(), im.size[0], im.size[1], locator)


def render_image(data: bytes, locator: str = "image") -> Rendered:
    from PIL import Image

    check_image_header(data)
    with Image.open(io.BytesIO(data)) as im:
        im.seek(0)  # a GIF or TIFF: the first frame
        frame = im.copy()
    return _encode(frame, locator)


def pdf_page_count(data: bytes) -> int:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(data)
    try:
        return len(doc)
    finally:
        doc.close()


def render_pdf_page(data: bytes, page: int) -> Rendered:
    """1-based page → JPEG. Forms and JavaScript aren't run (pdfium renders the page content)."""
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(data)
    try:
        if not 1 <= page <= len(doc):
            raise ValueError(f"The PDF has {len(doc)} pages")
        p = doc[page - 1]
        w_pt, h_pt = p.get_size()
        scale = DPI / 72
        if (w_pt * scale) * (h_pt * scale) > MAX_IMAGE_PIXELS:
            raise UnsafeFile("The page is too large to render")
        bitmap = p.render(scale=scale)
        im = bitmap.to_pil()
        p.close()
    finally:
        doc.close()
    return _encode(im, f"p{page}")


def embedded_images(data: bytes, filename: str) -> list[tuple[str, bytes]]:
    """Pictures inside a .docx (in order) or a .pptx (``slide N``), as (locator, bytes)."""
    ext = filename.rsplit(".", 1)[-1].lower()
    out: list[tuple[str, bytes]] = []
    if ext in ("docx", "docm"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        for n, shape in enumerate(doc.inline_shapes, start=1):
            blips = shape._inline.xpath(".//a:blip/@r:embed")
            if blips:
                part = doc.part.related_parts.get(blips[0])
                if part is not None:
                    out.append((f"image {n}", part.blob))
    elif ext in ("pptx", "pptm"):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = Presentation(io.BytesIO(data))
        for n, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    out.append((f"slide {n}", shape.image.blob))
    return out
