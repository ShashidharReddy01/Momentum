"""PowerPoint .pptx (python-pptx): slides as the outline (their titles), text frames, tables,
speaker notes and pictures, each located as ``slide N`` (notes: ``slide N notes``)."""

from __future__ import annotations

import io

from momentum.files.model import ParseResult
from momentum.files.parsers.base import Builder

EMU_PER_PX = 9525


def parse_pptx(data: bytes, filename: str, mime: str) -> ParseResult:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    b = Builder(data, filename, mime, "presentation")
    prs = Presentation(io.BytesIO(data))
    slides = list(prs.slides)
    b.header.slides = len(slides)
    for n, slide in enumerate(slides, start=1):
        loc = f"slide {n}"
        title_shape = slide.shapes.title
        title = (
            title_shape.text_frame.text.strip()
            if title_shape is not None and title_shape.has_text_frame
            else ""
        )
        b.heading(title or f"Slide {n}", 1, loc)
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                b.text(shape.text_frame.text, loc)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                b.table(rows, f"{loc} table", title=title or None)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                w = int(shape.width or 0) // EMU_PER_PX
                h = int(shape.height or 0) // EMU_PER_PX
                b.image(loc, width=w or None, height=h or None, slide=n, alt=shape.name)
        if slide.has_notes_slide:
            b.text(slide.notes_slide.notes_text_frame.text, f"{loc} notes", style="notes")
    return b.done()
