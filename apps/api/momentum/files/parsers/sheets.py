"""Spreadsheets: .xlsx/.xlsm/.xltx (openpyxl), .xls (xlrd), .csv/.tsv (stdlib).

Every sheet gets a header guess, typed columns, its rows (cached values; never recalculated),
formulas as text beside their cached values, named ranges, flattened merged cells, a hidden flag
and a chart count. Rows are kept apart from the model (``ParseResult.rows``) from cell A1, capped
at ``max_rows`` rows and 100 columns per sheet.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable

from momentum.files.macros import read_macros
from momentum.files.model import (
    Cell,
    FormulaCell,
    ParseResult,
    SheetColumn,
    SheetInfo,
    cell_value,
    column_letter,
)
from momentum.files.parsers.base import Builder
from momentum.files.safety import MAX_COLUMNS
from momentum.files.values import infer_type, is_empty

MAX_FORMULAS_LISTED = 500
HEADER_SCAN_ROWS = 10


def guess_header(rows: list[list[Cell]]) -> int:
    """1-based row of the header: the first of the first 10 rows with at least two filled cells,
    mostly text, followed by another filled row. 0 when there's no such row."""
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        filled = [c for c in row if not is_empty(c)]
        if len(filled) < 2:
            continue
        texty = sum(
            1 for c in filled if isinstance(c, str) and not c.strip().replace(".", "").isdigit()
        )
        if texty >= 0.6 * len(filled) and any(
            not is_empty(c) for r in rows[i + 1 : i + 3] for c in r
        ):
            return i + 1
        return 0  # the first substantial row isn't a header: treat the sheet as headerless
    return 0


def sheet_info(
    name: str, index: int, rows: list[list[Cell]], *, truncated: bool, **extra: object
) -> SheetInfo:
    width = min(max((len(r) for r in rows), default=0), MAX_COLUMNS)
    header = guess_header(rows)
    body = rows[header:] if header else rows
    columns: list[SheetColumn] = []
    head = rows[header - 1] if header else []
    for c in range(width):
        label = head[c] if c < len(head) and not is_empty(head[c]) else None
        letter = column_letter(c)
        columns.append(
            SheetColumn(
                name=str(label).strip() if label is not None else letter,
                letter=letter,
                inferred_type=infer_type([r[c] if c < len(r) else None for r in body]),
            )
        )
    data_rows = sum(1 for r in body if any(not is_empty(c) for c in r))
    dims = f"A1:{column_letter(max(width - 1, 0))}{len(rows)}" if rows else "A1:A1"
    return SheetInfo(
        name=name,
        dims=dims,
        header_row_guess=header,
        columns=columns,
        row_count=data_rows,
        truncated=truncated,
        data_ref=f"sheet{index}",
        **extra,
    )


def _grid(values: Iterable[Iterable[object]], max_rows: int) -> tuple[list[list[Cell]], bool]:
    out: list[list[Cell]] = []
    truncated = False
    for row in values:
        if len(out) >= max_rows:
            truncated = True
            break
        cells = [cell_value(v) for v in row]
        if len(cells) > MAX_COLUMNS:
            cells, truncated = cells[:MAX_COLUMNS], True
        out.append(cells)
    while out and all(is_empty(c) for c in out[-1]):
        out.pop()
    return out, truncated


def parse_xlsx(data: bytes, filename: str, mime: str, *, max_rows: int) -> ParseResult:
    from openpyxl import load_workbook

    b = Builder(data, filename, mime, "spreadsheet")
    values_wb = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    formulas_wb = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    names_by_sheet: dict[str, list[str]] = {}
    try:
        defined = list(formulas_wb.defined_names.items())
    except AttributeError:  # pragma: no cover - older openpyxl
        defined = []
    for nm, dn in defined:
        text = str(getattr(dn, "attr_text", "") or "")
        sheet = text.split("!")[0].strip("'") if "!" in text else ""
        names_by_sheet.setdefault(sheet, []).append(f"{nm} = {text}")
    b.header.sheets = len(values_wb.worksheets)
    no_values = 0
    for i, ws in enumerate(values_wb.worksheets):
        fws = formulas_wb[ws.title]
        rows, truncated = _grid(ws.iter_rows(values_only=True), max_rows)
        merged_set = getattr(ws, "merged_cells", None)
        merged = list(merged_set.ranges) if merged_set is not None else []
        for rng in merged:  # flatten: every cell of a merged range shows the top-left value
            top = (
                rows[rng.min_row - 1][rng.min_col - 1]
                if rng.min_row <= len(rows) and rng.min_col <= len(rows[rng.min_row - 1])
                else None
            )
            for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
                row = rows[r]
                row.extend([None] * (rng.max_col - len(row)))
                for c in range(rng.min_col - 1, min(rng.max_col, MAX_COLUMNS)):
                    row[c] = top
        formulas: list[FormulaCell] = []
        count = 0
        missing = 0
        for frow in fws.iter_rows(max_row=min(fws.max_row, max_rows)):
            for cell in frow:
                v = cell.value
                # an array formula is an object with .text; a plain one is the string itself
                formula = v if isinstance(v, str) else getattr(v, "text", None)
                if isinstance(formula, str) and formula.startswith("="):
                    count += 1
                    r, c = cell.row - 1, cell.column - 1
                    cached = rows[r][c] if r < len(rows) and c < len(rows[r]) else None
                    if cached is None:
                        missing += 1
                    if len(formulas) < MAX_FORMULAS_LISTED:
                        formulas.append(
                            FormulaCell(cell=cell.coordinate, formula=formula, cached=cached)
                        )
        no_values += missing
        info = sheet_info(
            ws.title,
            i,
            rows,
            truncated=truncated,
            has_formulas=count > 0,
            formulas=formulas,
            formulas_without_values=missing,
            named_ranges=names_by_sheet.get(ws.title, []),
            charts_count=len(getattr(fws, "_charts", []) or []),
            hidden=ws.sheet_state != "visible",
            merged_cells=len(merged),
        )
        b.model.sheets.append(info)
        b.rows[info.data_ref] = rows
        b.heading(ws.title, 1, f"s:{ws.title}")
        if truncated:
            b.header.truncated = True
            b.warn(
                f"Sheet {ws.title} is larger than Mo reads ({max_rows:,} rows, 100 columns): "
                "the rest was cut"
            )
    if no_values:
        b.warn(
            f"{no_values} formulas have no saved value (the workbook was never recalculated in "
            "Excel): their results can't be read, only the formulas"
        )
    if b.model.sheets and any(s.hidden for s in b.model.sheets):
        b.warn("Some sheets are hidden")
    if filename.lower().endswith((".xlsm", ".xltm")) or "macroenabled" in mime.lower():
        b.model.macros = read_macros(data, filename)
    return b.done()


def parse_xls(data: bytes, filename: str, mime: str, *, max_rows: int) -> ParseResult:
    import xlrd

    b = Builder(data, filename, mime, "spreadsheet")
    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    b.header.sheets = book.nsheets
    for i in range(book.nsheets):
        sh = book.sheet_by_index(i)

        def row_values(r: int, sh: xlrd.sheet.Sheet = sh) -> list[object]:
            out: list[object] = []
            for c in range(min(sh.ncols, MAX_COLUMNS)):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    out.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    out.append(bool(cell.value))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    out.append(None)
                else:
                    out.append(cell.value)
            return out

        rows, truncated = _grid((row_values(r) for r in range(sh.nrows)), max_rows)
        truncated = truncated or sh.ncols > MAX_COLUMNS
        info = sheet_info(sh.name, i, rows, truncated=truncated, hidden=sh.visibility != 0)
        b.model.sheets.append(info)
        b.rows[info.data_ref] = rows
        b.heading(sh.name, 1, f"s:{sh.name}")
        book.unload_sheet(i)
    b.warn("Legacy .xls: values only (formulas aren't read from this format)")
    return b.done()


def decode_text(data: bytes) -> str:
    """UTF-16 with a BOM, else UTF-8 (BOM or not), else Windows-1252 (old Excel exports)."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def parse_csv(data: bytes, filename: str, mime: str, *, max_rows: int) -> ParseResult:
    b = Builder(data, filename, mime, "spreadsheet")
    text = decode_text(data)
    if filename.lower().endswith(".tsv"):
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows, truncated = _grid(reader, max_rows)
    name = filename.rsplit(".", 1)[0] or "Sheet1"
    info = sheet_info(name, 0, rows, truncated=truncated)
    b.header.sheets = 1
    b.model.sheets.append(info)
    b.rows[info.data_ref] = rows
    b.heading(name, 1, f"s:{name}")
    if truncated:
        b.header.truncated = True
        b.warn(
            f"The file is larger than Mo reads ({max_rows:,} rows, 100 columns): the rest was cut"
        )
    if any(
        isinstance(c, str) and c[:1] in ("=", "+", "-", "@") and len(c) > 1 and not c[1:2].isdigit()
        for r in rows
        for c in r
    ):
        b.warn("Some cells start with = + - or @: they are shown as text, never run as formulas")
    return b.done()
